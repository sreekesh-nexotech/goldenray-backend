"""``attendance/reports/*`` — one calendar-fill source, one attendance-% formula, unambiguous codes (A9), CSV/XLSX with
formula-injection escaping, PDF and > 5,000 rows through the documents render job."""

from __future__ import annotations

import csv
import datetime as dt
import io

import pytest
from openpyxl import load_workbook

from attendance.services import exporters, recompute, reports
from attendance.tests.conftest import at
from attendance.tests.factories import punch
from documents.models import RenderJob
from hr.tests.factories import HolidayFactory, LeaveRecordFactory

pytestmark = pytest.mark.django_db

BASE = "/api/v1/attendance/reports/"


@pytest.fixture
def month(world):
    """Asha: present (7th), late (8th), half day (9th), absent (10th), leave (11th); Onam on the 3rd for everyone."""
    punch(world.d1, "1", at(7, 9, 30))
    punch(world.d1, "1", at(7, 18, 30))
    punch(world.d1, "1", at(8, 9, 50))
    punch(world.d1, "1", at(8, 18, 30))
    punch(world.d1, "1", at(9, 10, 30))
    punch(world.d1, "1", at(9, 14, 30))
    LeaveRecordFactory(employee=world.asha, date_from=dt.date(2026, 9, 11), date_to=dt.date(2026, 9, 11), status="APPROVED")
    HolidayFactory(date=dt.date(2026, 9, 3), name="Onam", office=None)
    recompute.recompute(date_from=dt.date(2026, 9, 7), date_to=dt.date(2026, 9, 11), reason="test")  # the rest stays unstored
    return world


def get(client, name, **params):
    return client.get(f"{BASE}{name}/", params)


class TestPermissionsAndScope:
    @pytest.mark.parametrize("name", ["daily", "weekly", "monthly", "monthly-detail", "monthly-individual", "individual", "office"])
    def test_export_permission(self, api_client, staff_client, name, world):
        assert api_client.get(f"{BASE}{name}/").status_code == 401
        response = staff_client.get(f"{BASE}{name}/")  # Staff: attendance.view only
        assert response.status_code == 403 and response.json()["code"] == "permission_denied"

    def test_the_office_manager_sees_their_office(self, manager_client, world):
        rows = get(manager_client, "daily", day="2026-09-14").json()["rows"]
        assert {row["employee_code"] for row in rows} == {"E003", "M001"}
        assert get(manager_client, "monthly", year=2026, month=9, employee=str(world.asha.uid)).status_code == 404


class TestReports:
    def test_daily_is_calendar_filled(self, hr_client, month):
        body = get(hr_client, "daily", day="2026-09-09").json()
        assert body["report"] == "daily" and body["title"] == "Daily Attendance Report" and body["row_count"] == 4
        rows = {row["employee_code"]: row for row in body["rows"]}
        assert (rows["E001"]["status"], rows["E001"]["status_label"], rows["E001"]["first_in"]) == ("HD", "Half Day", "10:30")
        assert rows["E002"]["status"] == "A"  # stored; eSSL read stored rows only
        assert body["totals"]["Half Day"] == 1 and body["totals"]["Employees"] == 4
        weekly_off = get(hr_client, "daily", day="2026-09-06").json()  # never stored: filled
        assert {row["status"] for row in weekly_off["rows"]} == {"WO"}
        today = get(hr_client, "daily").json()  # today: pending, never absent (A7)
        assert {row["status"] for row in today["rows"]} == {""} and {row["status_label"] for row in today["rows"]} == {"Pending"}
        only_absent = get(hr_client, "daily", day="2026-09-09", status="ABSENT").json()
        assert {row["employee_code"] for row in only_absent["rows"]} == {"E002", "E003", "M001"}

    def test_weekly_uses_unambiguous_codes(self, hr_client, month):
        body = get(hr_client, "weekly", week_start="2026-09-07").json()
        asha = next(row for row in body["rows"] if row["employee_code"] == "E001")
        assert [asha[f"2026-09-{day:02d}"] for day in range(7, 14)] == ["8:00", "7:40", "4:00", "A", "L", "A", "WO"]
        assert (asha["present_days"], asha["absent_days"], asha["half_day"], asha["late"]) == (3, 2, 1, 1)
        headers = [column["header"] for column in body["columns"]]
        assert "Mon 07" in headers and "Attendance %" in headers
        holiday_week = get(hr_client, "weekly", week_start="2026-08-31").json()
        binu = next(row for row in holiday_week["rows"] if row["employee_code"] == "E002")
        assert binu["2026-09-03"] == "H" and binu["2026-09-06"] == "WO"  # H is only ever a holiday (A9)

    def test_monthly_and_one_formula(self, hr_client, month):
        body = get(hr_client, "monthly", year=2026, month=9).json()
        asha = next(row for row in body["rows"] if row["employee_code"] == "E001")
        # present 1, late 1, half 1, leave 1, holiday 1, absent: working days of 1..14 without a punch
        assert (asha["present"], asha["late"], asha["half_day"], asha["leave"], asha["holiday"]) == (1, 1, 1, 1, 1)
        expected = asha["present"] + asha["late"] + asha["half_day"] + asha["absent"]
        assert asha["attendance_pct"] == f"{round((2 + 0.5) / expected * 100, 1)}%"
        assert body["totals"]["Days"] == 30

    def test_details_and_individual(self, hr_client, month):
        detail = get(hr_client, "monthly-detail", year=2026, month=9).json()
        assert detail["row_count"] == 4 * 30
        single = get(hr_client, "monthly-individual", year=2026, month=9, employee=str(month.asha.uid)).json()
        assert single["row_count"] == 30 and "device PINs 1" in single["subtitle"] and single["totals"]["Leave"] == 1
        everyone = get(hr_client, "monthly-individual", year=2026, month=9).json()
        assert everyone["report"] == "monthly-detail"
        window = get(hr_client, "individual", date_from="2026-09-07", date_to="2026-09-09", employee=str(month.asha.uid)).json()
        assert [row["status"] for row in window["rows"]] == ["P", "LT", "HD"] and "employee_code" not in window["columns"][0]["key"]
        all_people = get(hr_client, "individual", date_from="2026-09-07", date_to="2026-09-07").json()
        assert all_people["row_count"] == 4 and all_people["columns"][0]["key"] == "employee_code"

    def test_office_report_pools_people_with_the_same_formula(self, hr_client, month):
        body = get(hr_client, "office", date_from="2026-09-07", date_to="2026-09-11").json()
        rows = {row["office_code"]: row for row in body["rows"]}
        ho = rows["HO"]
        # HO: Asha P LT HD A L, Binu A A A A A (her leave is out of the denominator)
        assert (ho["employees"], ho["expected_days"], ho["present"], ho["half_day"], ho["absent"], ho["leave"]) == (2, 9, 2, 1, 6, 1)
        assert ho["attendance_pct"] == f"{round(2.5 / 9 * 100, 1)}%"
        assert body["totals"]["Attendance %"] == f"{round(2.5 / 19 * 100, 1)}%"  # overall: the same formula over every office

    def test_validation(self, hr_client, world):
        response = get(hr_client, "individual", date_from="2026-09-10", date_to="2026-09-01")
        assert response.status_code == 400
        response = get(hr_client, "office", date_from="2024-01-01", date_to="2026-09-01")
        assert response.status_code == 400 and response.json()["code"] == "range_too_long"
        assert get(hr_client, "monthly", year=2026, month=0).status_code == 400
        assert get(hr_client, "weekly").status_code == 400
        assert get(hr_client, "daily", format="docx").status_code == 400


class TestFiles:
    def test_csv_has_a_bom_a_summary_and_no_formulas(self, hr_client, month):
        month.binu.full_name = '=HYPERLINK("http://evil")'
        month.binu.save()
        response = get(hr_client, "daily", day="2026-09-09", format="csv")
        assert response.status_code == 200 and response["Content-Type"].startswith("text/csv")
        assert response["Content-Disposition"] == 'attachment; filename="daily-attendance-2026-09-09.csv"'
        content = response.content
        assert content.startswith(b"\xef\xbb\xbf")
        rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
        assert rows[0] == ["Daily Attendance Report"] and rows[3][0] == "Employee ID"
        assert any(cell == '\'=HYPERLINK("http://evil")' for row in rows for cell in row)
        assert ["Summary"] in rows

    def test_xlsx_opens_and_escapes(self, hr_client, month):
        month.binu.full_name = "@SUM(1+1)"
        month.binu.save()
        response = get(hr_client, "monthly", year=2026, month=9, format="xlsx")
        assert response.status_code == 200 and response["Content-Disposition"].endswith('monthly-attendance-2026-09.xlsx"')
        sheet = load_workbook(io.BytesIO(response.content)).active
        assert sheet["A1"].value == "Monthly Attendance Report" and sheet["A4"].value == "Employee ID"
        names = [sheet.cell(row=row, column=2).value for row in range(5, 9)]
        assert "'@SUM(1+1)" in names and all(cell.data_type != "f" for row in sheet.iter_rows() for cell in row)

    def test_safe_leaves_numbers_alone(self):
        assert exporters.safe(-5) == -5 and exporters.safe("-5") == "'-5" and exporters.safe("Asha") == "Asha" and exporters.safe(None) is None

    def test_pdf_is_a_render_job(self, hr_client, month, django_capture_on_commit_callbacks):
        with django_capture_on_commit_callbacks(execute=True):
            response = get(hr_client, "office", date_from="2026-09-07", date_to="2026-09-11", format="pdf")
        assert response.status_code == 202
        body = response.json()
        assert body["format"] == body["requested_format"] == "pdf" and body["job"]["kind"] == "ATTENDANCE_REPORT"
        job = RenderJob.objects.get(uid=body["job"]["uid"])
        assert job.object_type == "attendance.report" and job.payload["title"] == "Office Attendance Report" and job.status == "DONE"
        status = hr_client.get(f"/api/v1/documents/jobs/{job.uid}/")
        assert status.status_code == 200 and status.json()["status"] == "DONE"

    def test_large_reports_go_async(self, hr_client, month, monkeypatch):
        monkeypatch.setattr(exporters, "ASYNC_ROW_LIMIT", 10)
        for fmt in ("csv", "json"):
            response = get(hr_client, "monthly-detail", year=2026, month=9, format=fmt)
            assert response.status_code == 202 and response.json()["rows"] == 120 and response.json()["requested_format"] == fmt
        assert get(hr_client, "monthly", year=2026, month=9, format="csv").status_code == 200  # 4 rows

    def test_the_documents_view_follows_the_export_permission(self, hr_client, staff_client, month):
        body = get(hr_client, "daily", day="2026-09-09", format="pdf").json()
        assert staff_client.get(f"/api/v1/documents/jobs/{body['job']['uid']}/").status_code == 403

    def test_report_objects(self, month):
        report = reports.daily([month.asha], dt.date(2026, 9, 7))
        assert report.headers[0] == "Employee ID" and report.matrix()[0][0] == "E001"
