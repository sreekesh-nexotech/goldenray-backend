"""Report files: CSV and XLSX in the request; PDF — and any report over 5,000 rows — through a documents render job.

* **CSV** (UTF-8 with BOM for Excel): title, subtitle, a blank row, the header, the rows, then a ``Summary`` block.
  **XLSX** (openpyxl): title and subtitle merged, header on row 4, frozen panes, bounded column widths, summary.
* **Formula injection** (eSSL §I.7): every text cell starting with ``=``, ``+``, ``-``, ``@``, a tab or a carriage
  return is prefixed with ``'`` in both formats — names and device labels come from terminals and people. Numbers
  stay numbers.
* **PDF** is rendered only by the documents worker (Playwright; reportlab is gone, PLAN §1.4), so a PDF request is
  answered ``202`` with the ``ATTENDANCE_REPORT`` render job (``documents/jobs/<uid>/`` → ``download-url/``).
  A report of more than :data:`ASYNC_ROW_LIMIT` rows is rendered the same way whatever format was asked (DV-89): the
  request never holds a worker for a large export, and the frozen payload is capped at 1 MB by the documents package
  (narrow the selection beyond that).
"""

from __future__ import annotations

import csv
import io
import uuid

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from attendance.services.reports import Report

ASYNC_ROW_LIMIT = 5000
ACCENT = "1F4E79"
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
MEDIA_TYPES = {"csv": "text/csv; charset=utf-8", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
OBJECT_TYPE = "attendance.report"


def safe(value):
    """A cell value a spreadsheet will never evaluate."""
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def to_csv(report: Report) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow([safe(report.title)])
    writer.writerow([safe(report.subtitle)])
    writer.writerow([])
    writer.writerow([safe(header) for header in report.headers])
    for row in report.matrix():
        writer.writerow([safe(value) for value in row])
    if report.totals:
        writer.writerow([])
        writer.writerow(["Summary"])
        for key, value in report.totals.items():
            writer.writerow([safe(key), safe(value)])
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(report: Report) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Report"
    width = max(1, len(report.columns))
    sheet.cell(row=1, column=1, value=safe(report.title)).font = Font(size=14, bold=True, color=ACCENT)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    sheet.cell(row=2, column=1, value=safe(report.subtitle)).font = Font(size=10, italic=True)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    header_row = 4
    fill = PatternFill("solid", fgColor=ACCENT)
    for index, header in enumerate(report.headers, start=1):
        cell = sheet.cell(row=header_row, column=index, value=safe(header))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for offset, values in enumerate(report.matrix(), start=header_row + 1):
        for index, value in enumerate(values, start=1):
            sheet.cell(row=offset, column=index, value=safe(value))
    for index, (key, header) in enumerate(report.columns, start=1):
        longest = max([len(str(header))] + [len(str(row.get(key, ""))) for row in report.rows])
        sheet.column_dimensions[get_column_letter(index)].width = min(40, max(10, longest + 2))
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    if report.totals:
        start = header_row + len(report.rows) + 3
        sheet.cell(row=start, column=1, value="Summary").font = Font(bold=True, color=ACCENT)
        for offset, (key, value) in enumerate(report.totals.items(), start=1):
            sheet.cell(row=start + offset, column=1, value=safe(key)).font = Font(bold=True)
            sheet.cell(row=start + offset, column=2, value=safe(value))
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


RENDERERS = {"csv": to_csv, "xlsx": to_xlsx}


def needs_render_job(report: Report, fmt: str) -> bool:
    return fmt == "pdf" or (fmt in RENDERERS and len(report.rows) > ASYNC_ROW_LIMIT) or (fmt == "json" and len(report.rows) > ASYNC_ROW_LIMIT)


def render_file(report: Report, fmt: str) -> tuple[bytes, str, str]:
    """``(bytes, media type, file name)`` for csv / xlsx."""
    return RENDERERS[fmt](report), MEDIA_TYPES[fmt], f"{report.filename}.{fmt}"


def payload(report: Report, *, requested_format: str) -> dict:
    """The frozen document payload (rows as lists: the payload is capped at 1 MB)."""
    return {
        "report": report.kind,
        "title": report.title,
        "subtitle": report.subtitle,
        "headers": report.headers,
        "rows": [[str(value) if value is not None else "" for value in row] for row in report.matrix()],
        "totals": [[str(key), str(value)] for key, value in report.totals.items()],
        "landscape": len(report.columns) > 8,
        "generated_at": timezone.now().isoformat(),
        "requested_format": requested_format,
        "filename": f"{report.filename}.pdf",
    }


def request_pdf(report: Report, *, user, requested_format: str):
    """Queue the report's PDF (``ATTENDANCE_REPORT``); the job belongs to a fresh report uid (no stored record)."""
    from documents.services.jobs import request_render

    return request_render("ATTENDANCE_REPORT", OBJECT_TYPE, uuid.uuid4(), "default", "en", payload(report, requested_format=requested_format), user)
